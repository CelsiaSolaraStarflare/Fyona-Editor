"""
High-level export helpers that convert layout payloads into various formats.

PDF output is delegated to pdf_export.render_layout_to_pdf. Raster formats
use raster_export.rasterize_layout to draw the canvas into Pillow images,
which are then encoded as PNG/JPEG or embedded into DOCX/PPTX containers.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from pdf_export import PdfExportError, PdfRenderResult, PdfRenderStats, render_layout_to_pdf
from raster_export import RasterPage, rasterize_layout

try:  # Optional dependency for Word exports
    from docx import Document  # type: ignore
    from docx.shared import Inches  # type: ignore
except ImportError as exc:  # pragma: no cover
    DOCX_AVAILABLE = False
    DOCX_IMPORT_ERROR = exc
else:  # pragma: no cover
    DOCX_AVAILABLE = True
    DOCX_IMPORT_ERROR = None

try:  # Optional dependency for PowerPoint exports
    from pptx import Presentation  # type: ignore
except ImportError as exc:  # pragma: no cover
    PPTX_AVAILABLE = False
    PPTX_IMPORT_ERROR = exc
else:  # pragma: no cover
    PPTX_AVAILABLE = True
    PPTX_IMPORT_ERROR = None


class ExportFormatError(PdfExportError):
    """Raised when a requested export format cannot be produced."""


@dataclass
class ExportPayload:
    data: bytes
    mimetype: str
    filename: str
    digest: str
    stats: PdfRenderStats
    meta: Dict[str, str]


IMAGE_FORMATS = {"png", "jpeg", "jpg"}
DOC_FORMATS = {"doc", "docx", "word"}
PPT_FORMATS = {"ppt", "pptx", "powerpoint"}


def export_layout(
    layout: Dict[str, Any],
    *,
    project_name: Optional[str] = None,
    asset_base: Optional[Path] = None,
    export_format: str = "pdf",
) -> ExportPayload:
    fmt = (export_format or "pdf").lower()
    pdf_result = render_layout_to_pdf(layout, project_name=project_name, asset_base=asset_base, persist=False, verify_lossless=True)

    if fmt == "pdf":
        return ExportPayload(
            data=pdf_result.data,
            mimetype="application/pdf",
            filename=_build_filename(project_name, "pdf"),
            digest=pdf_result.digest,
            stats=pdf_result.stats,
            meta={"format": "pdf"},
        )

    rasters: Optional[List[RasterPage]] = None

    if fmt in IMAGE_FORMATS:
        rasters = _ensure_rasters(rasters, layout, asset_base)
        return _package_images(rasters, fmt, project_name, pdf_result)

    if fmt in DOC_FORMATS:
        rasters = _ensure_rasters(rasters, layout, asset_base)
        data = _rasters_to_docx(rasters)
        return ExportPayload(
            data=data,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=_build_filename(project_name, "docx"),
            digest=pdf_result.digest,
            stats=pdf_result.stats,
            meta={"format": "docx"},
        )

    if fmt in PPT_FORMATS:
        rasters = _ensure_rasters(rasters, layout, asset_base)
        data = _rasters_to_pptx(rasters)
        return ExportPayload(
            data=data,
            mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename=_build_filename(project_name, "pptx"),
            digest=pdf_result.digest,
            stats=pdf_result.stats,
            meta={"format": "pptx"},
        )

    raise ExportFormatError(f"Unsupported export format '{fmt}'.")


def _build_filename(project: Optional[str], extension: str) -> str:
    base = project or "layout"
    return f"{base}-layout.{extension}"


def _ensure_rasters(
    cached: Optional[List[RasterPage]],
    layout: Dict[str, Any],
    asset_base: Optional[Path],
) -> List[RasterPage]:
    if cached is not None:
        return cached
    return rasterize_layout(layout, asset_base=asset_base)


def _package_images(
    rasters: List[RasterPage],
    fmt: str,
    project_name: Optional[str],
    pdf_result: PdfRenderResult,
) -> ExportPayload:
    ext = "png" if fmt == "png" else "jpg"
    mimetype = "image/png" if fmt == "png" else "image/jpeg"

    if len(rasters) == 1:
        binary = _encode_raster(rasters[0], fmt)
        return ExportPayload(
            data=binary,
            mimetype=mimetype,
            filename=_build_filename(project_name, ext),
            digest=pdf_result.digest,
            stats=pdf_result.stats,
            meta={"format": ext},
        )

    archive_name = _build_filename(project_name, f"{ext}-pages")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for page in rasters:
            page_name = f"{(project_name or 'layout')}-page-{page.index + 1:02d}.{ext}"
            bundle.writestr(page_name, _encode_raster(page, fmt))
    buffer.seek(0)
    return ExportPayload(
        data=buffer.read(),
        mimetype="application/zip",
        filename=f"{archive_name}.zip",
        digest=pdf_result.digest,
        stats=pdf_result.stats,
        meta={"format": f"{ext}-zip"},
    )


def _encode_raster(page: RasterPage, fmt: str) -> bytes:
    buffer = io.BytesIO()
    image = page.image.convert("RGB")
    fmt_key = "png" if fmt == "png" else "jpeg"
    if fmt_key == "png":
        image.save(buffer, format="PNG", optimize=True)
    else:
        image.save(buffer, format="JPEG", quality=95, optimize=True)
    buffer.seek(0)
    return buffer.read()


def _require_docx():
    if not DOCX_AVAILABLE:
        raise ExportFormatError("DOCX export requires python-docx.") from DOCX_IMPORT_ERROR


def _require_pptx():
    if not PPTX_AVAILABLE:
        raise ExportFormatError("PowerPoint export requires python-pptx.") from PPTX_IMPORT_ERROR


def _rasters_to_docx(rasters: List[RasterPage]) -> bytes:
    _require_docx()
    document = Document()
    for index, page in enumerate(rasters):
        if index > 0:
            document.add_page_break()
        stream = io.BytesIO(_encode_raster(page, "png"))
        width_inches = max(page.width / 96.0, 1.0)
        document.add_picture(stream, width=Inches(width_inches))
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _rasters_to_pptx(rasters: List[RasterPage]) -> bytes:
    _require_pptx()
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    for page in rasters:
        slide = presentation.slides.add_slide(blank_layout)
        stream = io.BytesIO(_encode_raster(page, "png"))
        slide.shapes.add_picture(stream, 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return buffer.read()


__all__ = ["ExportFormatError", "ExportPayload", "export_layout"]
